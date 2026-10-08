#!/usr/bin/env python3
import argparse, csv, json, math, re, sqlite3
from pathlib import Path
if __package__:
    from .food_search_index import build_search_index
    from .meal_recommendations import build_meal_recommendations
else:
    from food_search_index import build_search_index
    from meal_recommendations import build_meal_recommendations

KJ_PER_KCAL = 4.184
SCHEMA = '''
CREATE TABLE foods (
  code TEXT PRIMARY KEY,
  barcode TEXT,
  name TEXT NOT NULL,
  search_name TEXT NOT NULL,
  brand TEXT,
  quantity TEXT,
  category TEXT,
  source TEXT NOT NULL,
  calories_100g REAL,
  protein_100g REAL,
  carbs_100g REAL,
  fat_100g REAL,
  sugars_100g REAL,
  fiber_100g REAL,
  saturated_fat_100g REAL,
  salt_100g REAL,
  sodium_100g REAL,
  calcium_mg_100g REAL,
  phosphorus_mg_100g REAL,
  magnesium_mg_100g REAL,
  sodium_mg_100g REAL,
  potassium_mg_100g REAL,
  iron_mg_100g REAL,
  copper_mg_100g REAL,
  selenium_ug_100g REAL,
  chromium_ug_100g REAL,
  manganese_mg_100g REAL,
  zinc_mg_100g REAL,
  vitamin_a_ug_100g REAL,
  vitamin_e_mg_100g REAL,
  vitamin_d_ug_100g REAL,
  vitamin_k_ug_100g REAL,
  thiamine_mg_100g REAL,
  riboflavin_mg_100g REAL,
  niacin_mg_100g REAL,
  pantothenic_acid_mg_100g REAL,
  vitamin_b6_mg_100g REAL,
  biotin_ug_100g REAL,
  folate_ug_100g REAL,
  vitamin_c_mg_100g REAL,
  serving_size TEXT,
  serving_g REAL,
  ingredients TEXT,
  allergens TEXT,
  labels TEXT,
  nutriscore TEXT,
  nova INTEGER,
  product_url TEXT,
  updated_at TEXT,
  extra_json TEXT
);
CREATE UNIQUE INDEX idx_foods_barcode ON foods(barcode) WHERE barcode IS NOT NULL AND barcode <> '';
CREATE INDEX idx_foods_search_name ON foods(search_name);
CREATE INDEX idx_foods_brand ON foods(brand);
CREATE INDEX idx_foods_source ON foods(source);
CREATE TABLE metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL);
'''

MICRO_COLS = [
    'calcium_mg_100g','phosphorus_mg_100g','magnesium_mg_100g','sodium_mg_100g',
    'potassium_mg_100g','iron_mg_100g','copper_mg_100g','selenium_ug_100g',
    'chromium_ug_100g','manganese_mg_100g','zinc_mg_100g','vitamin_a_ug_100g',
    'vitamin_e_mg_100g','vitamin_d_ug_100g','vitamin_k_ug_100g','thiamine_mg_100g',
    'riboflavin_mg_100g','niacin_mg_100g','pantothenic_acid_mg_100g','vitamin_b6_mg_100g',
    'biotin_ug_100g','folate_ug_100g','vitamin_c_mg_100g'
]

COLS = [
    'code','barcode','name','search_name','brand','quantity','category','source',
    'calories_100g','protein_100g','carbs_100g','fat_100g','sugars_100g','fiber_100g',
    'saturated_fat_100g','salt_100g','sodium_100g', *MICRO_COLS,
    'serving_size','serving_g','ingredients','allergens','labels','nutriscore','nova',
    'product_url','updated_at','extra_json'
]

def text(v):
    if v is None: return None
    s = str(v).strip()
    return s if s else None

def num(v):
    if v is None or v == '': return None
    try:
        n = float(v)
        return n if math.isfinite(n) else None
    except (TypeError, ValueError):
        return None

def scaled_num(row, key, factor):
    value = num(row.get(key))
    return value * factor if value is not None else None

def sum_known(*values):
    known = [v for v in values if v is not None]
    return sum(known) if known else None

def sodium_from_salt_g(salt_g):
    salt = num(salt_g)
    return salt / 2.5 * 1000.0 if salt is not None else None

def ifct_micros(row):
    return {
        'calcium_mg_100g': scaled_num(row, 'ca', 1000.0),
        'phosphorus_mg_100g': scaled_num(row, 'p', 1000.0),
        'magnesium_mg_100g': scaled_num(row, 'mg', 1000.0),
        'sodium_mg_100g': scaled_num(row, 'na', 1000.0),
        'potassium_mg_100g': scaled_num(row, 'k', 1000.0),
        'iron_mg_100g': scaled_num(row, 'fe', 1000.0),
        'copper_mg_100g': scaled_num(row, 'cu', 1000.0),
        'selenium_ug_100g': scaled_num(row, 'se', 1000000.0),
        'chromium_ug_100g': scaled_num(row, 'cr', 1000000.0),
        'manganese_mg_100g': scaled_num(row, 'mn', 1000.0),
        'zinc_mg_100g': scaled_num(row, 'zn', 1000.0),
        'vitamin_a_ug_100g': scaled_num(row, 'vita', 1000000.0),
        'vitamin_e_mg_100g': scaled_num(row, 'vite', 1000.0),
        'vitamin_d_ug_100g': scaled_num(row, 'vitd', 1000000.0),
        'vitamin_k_ug_100g': scaled_num(row, 'vitk', 1000000.0),
        'thiamine_mg_100g': scaled_num(row, 'thia', 1000.0),
        'riboflavin_mg_100g': scaled_num(row, 'ribf', 1000.0),
        'niacin_mg_100g': scaled_num(row, 'nia', 1000.0),
        'pantothenic_acid_mg_100g': scaled_num(row, 'pantac', 1000.0),
        'vitamin_b6_mg_100g': scaled_num(row, 'vitb6c', 1000.0),
        'biotin_ug_100g': scaled_num(row, 'biot', 1000000.0),
        'folate_ug_100g': scaled_num(row, 'folsum', 1000000.0),
        'vitamin_c_mg_100g': scaled_num(row, 'vitc', 1000.0),
    }

def indb_micros(row):
    vitamin_d = sum_known(num(row.get('vitd2_ug')), num(row.get('vitd3_ug')))
    vitamin_k = sum_known(num(row.get('vitk1_ug')), num(row.get('vitk2_ug')))
    folate = num(row.get('folate_ug'))
    if folate is None:
        folate = num(row.get('vitb9_ug'))
    return {
        'calcium_mg_100g': num(row.get('calcium_mg')),
        'phosphorus_mg_100g': num(row.get('phosphorus_mg')),
        'magnesium_mg_100g': num(row.get('magnesium_mg')),
        'sodium_mg_100g': num(row.get('sodium_mg')),
        'potassium_mg_100g': num(row.get('potassium_mg')),
        'iron_mg_100g': num(row.get('iron_mg')),
        'copper_mg_100g': num(row.get('copper_mg')),
        'selenium_ug_100g': num(row.get('selenium_ug')),
        'chromium_ug_100g': scaled_num(row, 'chromium_mg', 1000.0),
        'manganese_mg_100g': num(row.get('manganese_mg')),
        'zinc_mg_100g': num(row.get('zinc_mg')),
        'vitamin_a_ug_100g': num(row.get('vita_ug')),
        'vitamin_e_mg_100g': num(row.get('vite_mg')),
        'vitamin_d_ug_100g': vitamin_d,
        'vitamin_k_ug_100g': vitamin_k,
        'thiamine_mg_100g': num(row.get('vitb1_mg')),
        'riboflavin_mg_100g': num(row.get('vitb2_mg')),
        'niacin_mg_100g': num(row.get('vitb3_mg')),
        'pantothenic_acid_mg_100g': num(row.get('vitb5_mg')),
        'vitamin_b6_mg_100g': num(row.get('vitb6_mg')),
        'biotin_ug_100g': num(row.get('vitb7_ug')),
        'folate_ug_100g': folate,
        'vitamin_c_mg_100g': num(row.get('vitc_mg')),
    }

def integer(v):
    n = num(v)
    return int(n) if n is not None else None

def norm(s):
    return ' '.join((s or '').lower().split())

def payload(**kw):
    row = {k: None for k in COLS}
    row.update(kw)
    row['name'] = text(row['name']) or 'Unnamed food'
    row['search_name'] = norm(' '.join(filter(None, [row['name'], row.get('brand'), row.get('category')])))
    return row

def insert(conn, row):
    sql = f"INSERT OR REPLACE INTO foods ({','.join(COLS)}) VALUES ({','.join('?' for _ in COLS)})"
    conn.execute(sql, [row.get(c) for c in COLS])

def merge_json(old, key, obj):
    try: base = json.loads(old) if old else {}
    except Exception: base = {}
    base[key] = obj
    return json.dumps(base, ensure_ascii=False, separators=(',', ':'))

def import_off_csv(conn, path):
    count = 0
    with open(path, encoding='utf-8-sig', newline='') as f:
        for r in csv.DictReader(f):
            barcode = text(r.get('code'))
            name = text(r.get('product_name'))
            if not barcode or not name: continue
            insert(conn, payload(
                code='OFF-' + barcode, barcode=barcode, name=name,
                brand=text(r.get('brands')), category=text(r.get('categories')), source='OFF',
                calories_100g=num(r.get('energy-kcal_100g')), protein_100g=num(r.get('proteins_100g')),
                carbs_100g=num(r.get('carbohydrates_100g')), fat_100g=num(r.get('fat_100g')),
                sugars_100g=num(r.get('sugars_100g')), fiber_100g=num(r.get('fiber_100g')),
                saturated_fat_100g=num(r.get('saturated-fat_100g')), salt_100g=num(r.get('salt_100g')),
                sodium_100g=(num(r.get('salt_100g')) / 2.5 if num(r.get('salt_100g')) is not None else None),
                sodium_mg_100g=sodium_from_salt_g(r.get('salt_100g')),
                nutriscore=text(r.get('nutriscore_grade')), nova=integer(r.get('nova_group')),
                product_url=text(r.get('url')), updated_at=text(r.get('last_modified_datetime')),
                extra_json=json.dumps({'off': r}, ensure_ascii=False, separators=(',', ':'))
            ))
            count += 1
    return count

def import_openlabel(conn, path):
    raw = Path(path).read_text(encoding='utf-8')
    m = re.search(r'export const OFF_INDIA_DATA = ("(?:\\.|[^"\\])*");', raw, re.S)
    if not m: raise RuntimeError('Could not locate OFF_INDIA_DATA in OpenLabel file')
    rows = json.loads(json.loads(m.group(1)))
    enriched = inserted = 0
    for r in rows:
        barcode = text(r.get('barcode'))
        if not barcode: continue
        existing = conn.execute('SELECT * FROM foods WHERE barcode=?', (barcode,)).fetchone()
        n = r.get('nutrition') or {}
        ingredients = r.get('ingredients') or []
        ingredients_text = '\n'.join(str(x) for x in ingredients if x)
        if existing:
            old = dict(existing)
            updates = {
                'brand': old['brand'] or text(r.get('brand')),
                'category': old['category'] or text(r.get('category')),
                'ingredients': old['ingredients'] or text(ingredients_text),
                'calories_100g': old['calories_100g'] if old['calories_100g'] is not None else num(n.get('energyKcal100g')),
                'protein_100g': old['protein_100g'] if old['protein_100g'] is not None else num(n.get('proteins100g')),
                'sugars_100g': old['sugars_100g'] if old['sugars_100g'] is not None else num(n.get('sugars100g')),
                'fiber_100g': old['fiber_100g'] if old['fiber_100g'] is not None else num(n.get('fiber100g')),
                'saturated_fat_100g': old['saturated_fat_100g'] if old['saturated_fat_100g'] is not None else num(n.get('saturatedFat100g')),
                'salt_100g': old['salt_100g'] if old['salt_100g'] is not None else num(n.get('salt100g')),
                'sodium_100g': old['sodium_100g'] if old['sodium_100g'] is not None else (num(n.get('salt100g')) / 2.5 if num(n.get('salt100g')) is not None else None),
                'sodium_mg_100g': old['sodium_mg_100g'] if old['sodium_mg_100g'] is not None else sodium_from_salt_g(n.get('salt100g')),
                'nova': old['nova'] if old['nova'] is not None else integer(n.get('novaGroup')),
                'extra_json': merge_json(old['extra_json'], 'openlabel', r),
            }
            updates['search_name'] = norm(' '.join(filter(None, [old['name'], updates['brand'], updates['category']])))
            conn.execute('''UPDATE foods SET brand=:brand,category=:category,ingredients=:ingredients,
                calories_100g=:calories_100g,protein_100g=:protein_100g,sugars_100g=:sugars_100g,
                fiber_100g=:fiber_100g,saturated_fat_100g=:saturated_fat_100g,salt_100g=:salt_100g,
                sodium_100g=:sodium_100g,sodium_mg_100g=:sodium_mg_100g,
                nova=:nova,search_name=:search_name,extra_json=:extra_json WHERE barcode=:barcode''', {**updates, 'barcode': barcode})
            enriched += 1
        else:
            insert(conn, payload(
                code='OFF-' + barcode, barcode=barcode, name=text(r.get('name')) or 'Unnamed product',
                brand=text(r.get('brand')), category=text(r.get('category')), source='OFF',
                calories_100g=num(n.get('energyKcal100g')), protein_100g=num(n.get('proteins100g')),
                sugars_100g=num(n.get('sugars100g')), fiber_100g=num(n.get('fiber100g')),
                saturated_fat_100g=num(n.get('saturatedFat100g')), salt_100g=num(n.get('salt100g')),
                sodium_100g=(num(n.get('salt100g')) / 2.5 if num(n.get('salt100g')) is not None else None),
                sodium_mg_100g=sodium_from_salt_g(n.get('salt100g')),
                nova=integer(n.get('novaGroup')), ingredients=text(ingredients_text),
                extra_json=json.dumps({'openlabel': r}, ensure_ascii=False, separators=(',', ':'))
            ))
            inserted += 1
    return len(rows), enriched, inserted

def import_ifct(conn, path):
    count = 0
    with open(path, encoding='utf-8-sig', newline='') as f:
        for r in csv.DictReader(line for line in f if not line.startswith('#')):
            name = text(r.get('name')); code = text(r.get('code'))
            if not name or not code: continue
            protein, carbs, fat = num(r.get('protcnt')), num(r.get('choavldf')), num(r.get('fatce'))
            kj = num(r.get('enerc'))
            kcal = round(kj / KJ_PER_KCAL, 1) if kj is not None and kj > 0 else None
            if kcal is None and any(v is not None for v in (protein, carbs, fat)):
                kcal = round(4*(protein or 0) + 4*(carbs or 0) + 9*(fat or 0), 1)
            insert(conn, payload(
                code='IFCT-' + code, name=name, category=text(r.get('grup')), source='IFCT',
                calories_100g=kcal, protein_100g=protein, carbs_100g=carbs, fat_100g=fat,
                fiber_100g=num(r.get('fibtg')), sodium_100g=num(r.get('na')),
                **ifct_micros(r),
                extra_json=json.dumps({'ifct': r}, ensure_ascii=False, separators=(',', ':'))
            )); count += 1
    return count

def import_flat(conn, path, source_override=None):
    count=0
    with open(path, encoding='utf-8-sig', newline='') as f:
        for r in csv.DictReader(f):
            name=text(r.get('name')); code=text(r.get('code'))
            if not name or not code: continue
            source=source_override or text(r.get('source')) or 'GENERIC'
            insert(conn, payload(
                code=code, name=name, category='Foods & drinks' if source=='USDA' else 'Indian recipes', source=source,
                calories_100g=num(r.get('cal')), protein_100g=num(r.get('protein_g')), carbs_100g=num(r.get('carb_g')),
                fat_100g=num(r.get('fat_g')), fiber_100g=num(r.get('fiber_g')), serving_size=text(r.get('serving_desc')),
                serving_g=num(r.get('serving_g')), extra_json=json.dumps({source.lower(): r}, ensure_ascii=False, separators=(',', ':'))
            )); count += 1
    return count

def import_indb_xlsx(conn, path):
    from openpyxl import load_workbook
    wb=load_workbook(path, read_only=True, data_only=True)
    ws=wb['Nutrient Data']
    it=ws.iter_rows(values_only=True)
    headers=[str(v).strip() if v is not None else '' for v in next(it)]
    count=0
    for vals in it:
        r={headers[i]: vals[i] for i in range(min(len(headers),len(vals)))}
        name=text(r.get('food_name')); code=text(r.get('food_code'))
        if not name or not code: continue
        kcal=num(r.get('energy_kcal')); serving_kcal=num(r.get('unit_serving_energy_kcal'))
        serving_g=round((serving_kcal/kcal)*100) if kcal and serving_kcal else None
        insert(conn, payload(
            code='INDB-'+code, name=name, category='Indian cooked dishes', source='INDB', calories_100g=kcal,
            protein_100g=num(r.get('protein_g')), carbs_100g=num(r.get('carb_g')), fat_100g=num(r.get('fat_g')),
            fiber_100g=num(r.get('fibre_g')), sodium_100g=(num(r.get('sodium_mg')) / 1000.0 if num(r.get('sodium_mg')) is not None else None),
            **indb_micros(r),
            serving_size=text(r.get('servings_unit')), serving_g=serving_g,
            extra_json=json.dumps({'indb': r}, ensure_ascii=False, default=str, separators=(',', ':'))
        )); count += 1
    return count

def validate(conn):
    total=conn.execute('SELECT COUNT(*) FROM foods').fetchone()[0]
    packaged=conn.execute("SELECT COUNT(*) FROM foods WHERE source='OFF'").fetchone()[0]
    macros=conn.execute('''SELECT COUNT(*) FROM foods WHERE calories_100g IS NOT NULL AND protein_100g IS NOT NULL
                           AND carbs_100g IS NOT NULL AND fat_100g IS NOT NULL''').fetchone()[0]
    if total < 1: raise RuntimeError('database has no foods')
    dup=conn.execute("SELECT COUNT(*) FROM (SELECT barcode FROM foods WHERE barcode IS NOT NULL GROUP BY barcode HAVING COUNT(*)>1)").fetchone()[0]
    if dup: raise RuntimeError(f'duplicate barcodes: {dup}')
    return total, packaged, macros

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--off'); ap.add_argument('--openlabel'); ap.add_argument('--ifct'); ap.add_argument('--usda'); ap.add_argument('--indb')
    ap.add_argument('--output', required=True)
    args=ap.parse_args()
    out=Path(args.output); out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists(): out.unlink()
    conn=sqlite3.connect(out); conn.row_factory=sqlite3.Row; conn.executescript(SCHEMA)
    stats={}
    if args.off: stats['off_csv']=import_off_csv(conn,args.off)
    if args.openlabel:
        n,e,i=import_openlabel(conn,args.openlabel); stats.update(openlabel_rows=n,openlabel_enriched=e,openlabel_inserted=i)
    if args.ifct: stats['ifct']=import_ifct(conn,args.ifct)
    if args.usda: stats['usda']=import_flat(conn,args.usda,'USDA')
    if args.indb: stats['indb']=import_indb_xlsx(conn,args.indb)
    stats['regional_foods'], stats['search_terms'] = build_search_index(conn)
    total,packaged,macros=validate(conn)
    stats['meal_suggestions'] = build_meal_recommendations(conn, strict=total >= 26000)
    stats.update(total=total,packaged=packaged,complete_core_macros=macros)
    for k,v in stats.items(): conn.execute('INSERT INTO metadata(key,value) VALUES(?,?)',(k,str(v)))
    conn.execute('INSERT INTO metadata(key,value) VALUES(?,?)',('schema_version','2'))
    conn.commit(); conn.execute('VACUUM'); conn.close()
    print(json.dumps(stats, indent=2)); print(f'bytes={out.stat().st_size}')

if __name__=='__main__': main()
